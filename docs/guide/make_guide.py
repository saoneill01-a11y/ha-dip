"""Builds docs/DIP-assembly-guide.docx from the photos, renders and diagrams in this folder.
Run make_renders.py and make_diagrams.py first. Needs python-docx and Pillow."""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

G = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(G, "..", "DIP-assembly-guide.docx")
BLUE = RGBColor(0x1F, 0x4E, 0x8C)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.9))
sec.top_margin = sec.bottom_margin = Inches(0.8)
st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
    h = doc.styles[name]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.color.rgb = BLUE
    h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)

# footer page numbers
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("DIP assembly guide  |  page "); r.font.size = Pt(9); r.font.color.rgb = GREY
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run(); run.font.size = Pt(9); run.font.color.rgb = GREY
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = text; run._r.append(e)


def H(text, lvl=1):
    return doc.add_heading(text, lvl)


def P(text="", bold_lead=None, italic=False, size=None, color=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p


def B(text, bold_lead=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def STEP(n, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.45); p.paragraph_format.first_line_indent = Inches(-0.45)
    r = p.add_run(f"{n}.\t"); r.bold = True; r.font.color.rgb = BLUE
    p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def IMG(f, w=6.0, cap=None):
    doc.add_picture(os.path.join(G, f), width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY


def IMG2(f1, f2, w=3.2, cap=None):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, f in zip(t.rows[0].cells, (f1, f2)):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(os.path.join(G, f), width=Inches(w))
    if cap:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY


def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcol); tcPr.append(s)


def TABLE(rows, widths, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = Inches(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(val); r.font.size = Pt(10)
            if header and i == 0:
                r.bold = True; shade(c, "DCE6F2")
    doc.add_paragraph()
    return t


def NOTE(text, lead="Tip: ", fill="FFF6DD"):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); c.width = Inches(6.6); shade(c, fill)
    p = c.paragraphs[0]; p.add_run(lead).bold = True; p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def WARN(text):
    NOTE(text, "Careful: ", "FDE2E1")


def BREAK():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ================================================================ cover
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("DIP"); r.bold = True; r.font.size = Pt(48); r.font.color.rgb = BLUE
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Depth In Pond: pond water-level sensor"); r.font.size = Pt(18)
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Assembly, soldering and testing guide"); r.font.size = Pt(13); r.font.color.rgb = GREY
IMG("r_front.png", 3.6)
P("This guide takes you from a pile of parts to a working sensor at the pond. You don't need to know "
  "electronics. Every connection is shown in a picture, and you test as you go, so a mistake gets caught "
  "on the bench, not out by the water.")
P("What DIP does: ", None)
B("An ultrasonic sensor looks down a 3\" PVC pipe standing in the pond and measures how far down the water is. "
  "The pipe keeps the waterfall and ripples from confusing it.")
B("A float switch in the same pipe wakes the sensor the instant the water drops too low.")
B("A water temperature probe hangs in the pipe.")
B("Every 20 minutes it wakes, sends its readings to Home Assistant over Wi-Fi, and goes back to sleep. "
  "Home Assistant graphs the level, turns off one pump when the water drops 3\", and turns off all the pumps "
  "with an urgent phone alert when the water drops below the float, so they don't run dry.")
B("It runs on one rechargeable 18650 battery for months.")
BREAK()

# ================================================================ contents
H("What's in this guide")
for s in ["1. Tools and supplies", "2. The parts", "3. Printing the case", "4. Soldering basics (read this first)",
          "5. Prepare the printed parts", "6. Check the float switch", "7. Wire the battery holder, SuperMini and sensor board",
          "8. Bench test", "9. The probe cables", "10. The 4-core cable and the pipe cap", "11. Put the box together",
          "12. Software: ESPHome and Home Assistant", "13. Install at the pond", "14. Final checks and troubleshooting",
          "Appendix: every connection in one table"]:
    p = doc.add_paragraph(s); p.paragraph_format.space_after = Pt(2)
P()
H("The whole circuit", 2)
P("This is the full wiring. Don't worry if it looks busy: chapter 7 builds it one wire at a time. Come back to "
  "this page whenever you want to see where a wire fits in the big picture.")
IMG("d_overview.png", 6.7, "Everything DIP connects to. Colors match the steps later in the guide.")
BREAK()

# ================================================================ 1 tools
H("1. Tools and supplies")
TABLE([["Tool / supply", "What it's for"],
       ["Soldering iron, fine tip (a temperature-controlled one at about 340 °C / 650 °F is ideal)", "Joining wires and fitting the heat-set inserts"],
       ["Solder: thin (0.6-0.8 mm) rosin-core, lead-free or 60/40", "The metal that makes the joints"],
       ["Wire strippers", "Taking the plastic off wire ends"],
       ["Flush cutters (small side cutters)", "Cutting wire and resistor legs"],
       ["Heat shrink tubing, assorted small sizes, plus some adhesive-lined (glue-lined) pieces", "Insulating every joint; the adhesive kind seals out water"],
       ["Heat gun (a lighter works, carefully)", "Shrinking the tubing"],
       ["Multimeter (any cheap one with a beep/continuity mode)", "Checking the float switch, the battery and your joints"],
       ["\"Helping hands\" clamp or a lump of Blu-Tack", "Holding parts still while you solder"],
       ["Stranded hookup wire, 22-26 AWG, in red, black and a few other colors", "Short wires inside the box"],
       ["4 female-to-female jumper wires", "Plugging onto the sensor board's 4 pins"],
       ["USB-C cable and a computer", "Loading the software onto the SuperMini"],
       ["Small screwdrivers, drill with 1/4\", 1/8\" and pilot bits, marker, tape", "Assembly and the PVC pipe"],
       ["Safety glasses, a fan or open window", "Solder can spit, and the smoke isn't good to breathe"]],
      [3.6, 3.0])

# ================================================================ 2 parts
H("2. The parts")
P("Lay everything out and check it off before you start.")
TABLE([["Part", "How many", "Notes"],
       ["ESP32-C3 SuperMini board", "1", "The small black board with a USB-C socket. This is the brain."],
       ["18650 battery holder board + 18650 cell", "1 + 1", "Charges through USB-C. Has three 3V and three 5V output pad pairs."],
       ["JSN-SR20 dual-probe ultrasonic sensor (sold as JSN-SR20-Y1) + 2 probes", "1 set", "Blue board; the two round black probes plug into it."],
       ["Float switch (vertical, M10 stem)", "1", "Two yellow wires, either way round."],
       ["DS18B20 waterproof temperature probe", "1", "Steel tip on a black cable, red / black / yellow wires. Comes with a 4.7k resistor."],
       ["4.7k resistor", "1", "Came with the temperature probe."],
       ["100k resistors", "3", "Called (A), (B) and (C) in this guide."],
       ["4-core outdoor cable", "box to pipe, plus slack", "Four wires inside one jacket."],
       ["PG7 cable glands", "3", "Waterproof cable entries for the box."],
       ["M3 heat-set inserts", "7", "4 for the box lid, 3 for the pipe cap top."],
       ["M3 screws: 4 with rubber washers (box lid), 3 more (pipe cap top)", "7", "About 10-12 mm long."],
       ["M2.5 x 6 screws", "5", "4 for the sensor board, 1 for the SuperMini clamp."],
       ["Screws for the battery holder", "4", "M3 x 6, or the ones that came with the holder."],
       ["M4 screws, stainless", "2", "Float bracket to the pipe."],
       ["Wood screws, stainless", "2", "Hang the box on the post through its top tab."],
       ["Short stainless screw", "1", "Locks the pipe cap to the pipe."],
       ["3\" PVC pipe, about 30\" long", "1", "89 mm outside, 76 mm inside."],
       ["Cinder block + a bag of pea gravel", "1", "Holds the pipe upright in the pond."]],
      [2.6, 1.0, 3.0])
H("What the electronic parts look like", 2)
IMG2("supermini_front.jpg", "supermini_back.jpg", 3.1, "ESP32-C3 SuperMini, front (left) and back (right). The pin numbers are printed on the BACK.")
IMG2("holder_top_b.jpg", "holder_back.jpg", 3.1, "Battery holder, top and back. The pad pairs along the edges are marked 3V on one side and 5V on the other. The square pad of each pair is +.")
IMG2("sensor_front.jpg", "sensor_back.jpg", 3.1, "Ultrasonic sensor board, parts side (left) and connector side (right). Note the empty MODE pad top left: leave it empty.")
IMG2("probes.jpg", "temp_probe.jpg", 3.1, "The two ultrasonic probes on their cables (left) and the temperature probe (right).")
IMG2("float_switch.jpg", "resistor_4k7.jpg", 3.1, "Float switch in its bracket (left) and the 4.7k resistor that came with the temperature probe (right).")
H("Telling the resistors apart", 2)
P("Resistors are tiny and their value is shown by colored stripes. You have two kinds:")
IMG("d_resistors.png", 6.2)
P("If your stripes are hard to read, set your multimeter to resistance (Ω) and touch the probes to the two legs: "
  "a 100k reads close to 100 kΩ, the 4.7k reads close to 4.7 kΩ.")
BREAK()

# ================================================================ 3 printing
H("3. Printing the case")
P("All the printed parts are in the case/out folder of the project as STL files. PETG or ASA is best outdoors; "
  "PLA goes soft in the summer sun.")
TABLE([["File", "What it is", "How to print"],
       ["box.stl", "The electronics box that hangs on the post", "As it comes, back plate down, no supports"],
       ["box_lid.stl", "The lid, with DIP lettering", "As it comes (already flipped)"],
       ["esp_clamp.stl", "Small tab that holds the SuperMini down", "Flat"],
       ["gasket.stl", "Rubber-like seal for the lid", "95A TPU, flat, 100% infill (or use 1.5 mm foam cord instead)"],
       ["pipe_cap.stl", "Sits on top of the PVC pipe and holds the probes", "Upside down (rim on the bed); supports only under the probe band"],
       ["pipe_cap_top.stl", "Flat lid for the pipe cap", "Flat"],
       ["float_bracket.stl", "Holds the float switch inside the pipe", "As it comes (lying on its side)"]],
      [1.6, 2.6, 2.4])
IMG("r_box.png", 6.4, "The box")
IMG("r_box_lid.png", 6.4, "The lid: outside and inside")
IMG("r_pipe_cap.png", 6.4, "The pipe cap: top side, and underneath where the probes push in")
IMG2("r_pipe_cap_top.png", "r_float_bracket.png", 3.0, "Pipe cap top (left) and float bracket (right)")
IMG2("r_esp_clamp.png", "r_gasket.png", 3.0, "SuperMini clamp tab (left) and gasket (right)")
BREAK()

# ================================================================ 4 soldering basics
H("4. Soldering basics (read this first)")
P("Soldering melts a little metal (solder) into a joint so it holds and conducts. Good joints are shiny and "
  "smooth, like a little volcano or a smooth sleeve. Dull, lumpy blobs are \"cold\" joints and should be "
  "re-heated.")
H("How to make a joint", 2)
STEP(1, "Heat the iron. Wipe the tip on a damp sponge or brass wool, then touch a bit of solder to it so the tip is shiny (\"tinning\").")
STEP(2, "Hold the tip against BOTH things you're joining (the wire and the pad) for 1-2 seconds.")
STEP(3, "Feed solder into the joint, not onto the iron. It should flow and wrap around in a second or two.")
STEP(4, "Take the solder away, then the iron. Don't move the joint for a few seconds while it cools.")
NOTE("Tin wire ends before joining them: strip, twist the strands, then melt a little solder into them. "
     "Tinned wires go into holes easily and join in a second.")
H("Stripping wire", 2)
P("Strip about 5 mm for wires going into a board hole and about 8 mm for wire-to-wire splices. Don't nick the copper strands.")
H("Splicing two wires (you'll do this a lot)", 2)
IMG("d_splice.png", 6.2)
WARN("Always slide the heat shrink onto the wire BEFORE you solder. Everyone forgets this once; it means cutting the joint and doing it again.")
H("Soldering into the SuperMini", 2)
P("The SuperMini's pins are half-holes along its edges. Push the tinned wire into the hole from the front, hold "
  "the iron against the gold pad and the wire, and feed in a little solder. Some holes take two wires (or a "
  "wire and a resistor leg): put both in, then solder them together.")
WARN("Use as little solder as it takes. Too much can bridge to the next pin. If two pins get joined, touch the iron to the "
     "bridge and drag it away, or use solder wick.")
BREAK()

# ================================================================ 5 prepare printed parts
H("5. Prepare the printed parts")
H("Heat-set inserts", 2)
P("Inserts are little brass threads you melt into the plastic so screws hold strongly.")
STEP(1, "Find the 4 holes in the corner posts of the box and the 3 holes in the pipe cap's posts.")
STEP(2, "Set an insert on top of a hole, narrow end down.")
STEP(3, "Touch the hot soldering iron (a spare tip if you have one) into the insert and let it sink straight down slowly. Don't push hard. Stop when it's flush with the surface.")
STEP(4, "While it's still soft, press it flat with something metal. Let it cool.")
H("Cable glands", 2)
P("Screw a PG7 gland into each of the 3 holes in the bottom end of the box, with its lock nut inside. The nuts just fit "
  "side by side: turn each one by hand, then snug it with pliers. Leave the domed caps loose for now: cables go through later.")
IMG("r_box_filled.png", 6.4, "How it all goes together inside the box (battery holder, sensor board and SuperMini), and the 3 gland holes in the bottom end")
H("Gasket", 2)
P("Press the TPU gasket into the groove around the rim of the box (or lay 1.5 mm foam cord in the groove).")

# ================================================================ 6 float
H("6. Check the float switch")
P("The float switch is just a switch with no polarity: either wire can go anywhere. What matters is which way up it is. "
  "It must be OPEN (not connected) when the water is normal, and CLOSED when the water drops.")
STEP(1, "Set the multimeter to continuity (the beep setting). Touch the probes to the two yellow wires.")
STEP(2, "Hold the float the way it will sit in the pipe, stem up, float pushed UP toward the stem as if floating on water. The meter should be SILENT.")
STEP(3, "Let the float drop down (the water is gone). The meter should BEEP.")
STEP(4, "If it's the other way round, slide the clip off the bottom of the stem, turn the float over, and put the clip back. Check again.")
STEP(5, "Mount the float in the printed bracket: stem through the hole, nut on top, snug but don't crack the plastic.")
BREAK()

# ================================================================ 7 wiring
H("7. Wire the battery holder, SuperMini and sensor board")
WARN("Take the battery OUT of the holder while you solder. Put it back only when a step tells you to.")
P("First, find your way around the SuperMini. Turn it over so you can read the numbers on the back:")
IMG("d_pins.png", 5.8)
P("You'll use 7 pins: 3.3, G, 3, 4 on one edge, and 5, 6, 7 on the other. The rest stay empty.")
H("7a. Battery reading resistors (A) and (B)", 2)
P("These two 100k resistors let the SuperMini measure how full the battery is.")
IMG("d_divider.png", 6.0)
STEP(1, "Slide a short piece of heat shrink over each leg of resistor (B). Bend the legs down so it bridges from hole G to hole 4, passing over hole 3.3 without touching it.")
STEP(2, "Push the legs into G and 4 from the front side. Don't solder yet.")
STEP(3, "Bend one leg of resistor (A) and push it into hole 4 next to (B)'s leg. Solder hole 4 (both legs) and hole G. Trim the extra leg length on the back with the cutters.")
STEP(4, "Solder a 10 cm wire (a purple or other spare color) onto the free leg of resistor (A), with heat shrink over the joint. This wire will go to the battery's + contact.")
H("7b. Power from the battery holder to the SuperMini", 2)
STEP(1, "Cut a red and a black wire, about 10 cm each. Strip and tin both ends.")
STEP(2, "Solder the red wire into hole 3.3 and the black wire into hole G (G already has a resistor leg: the wire goes in alongside it).")
STEP(3, "Solder the other ends to the holder's 3V pair 1: red to the square + pad, black to the round - pad.")
STEP(4, "Solder the purple wire from resistor (A) to the holder's + battery contact. This is the metal clip at the end of the battery bay marked +; solder to where its tab comes through the back of the board.")
NOTE("Check the + battery contact before you solder: put the battery in, set the meter to DC volts, and touch black to a 3V - pad and red to the contact. "
     "It should read between 3.0 and 4.2 V. Take the battery back out before soldering.")
H("7c. The sensor board", 2)
P("The sensor board has 4 pins sticking out, labelled 5V, Trig, Echo and GND (the labels are on the connector side). Use female jumper wires "
  "so you can plug on rather than solder to the board.")
STEP(1, "Take 4 female jumper wires. Cut each to about 12 cm, keeping the plug end. Strip and tin the cut ends.")
STEP(2, "Plug them onto the sensor board pins. Use red for 5V and black for GND if you have them.")
STEP(3, "Trig wire into SuperMini hole 6, Echo wire into hole 7. Solder.")
STEP(4, "5V wire to the holder's 3V pair 2 + (square pad). GND wire to 3V pair 2 - (round pad). Solder.")
NOTE("Yes, the sensor's \"5V\" pin goes to a 3V pad. This sensor normally works on 3.3V, and running it on 3.3V keeps its signals safe for the SuperMini. "
     "The bench test in chapter 8 confirms it on your board.", "Why 3V? ")
WARN("Leave the MODE pad on the sensor board (two little pads top left, next to the word Mode) EMPTY. Putting a resistor there changes how the sensor talks and the readings stop.")
BREAK()

# ================================================================ 8 bench test
H("8. Bench test")
P("Now test what you've built before anything goes outside. The software has to be loaded first: chapter 12 explains how, "
  "but if someone else is handling the software, give them the SuperMini at this point (still wired up).")
STEP(1, "Plug the two ultrasonic probes into the white sockets on the sensor board (they only fit one way round).")
STEP(2, "Put the battery in. Slide the holder's switch to the \"stay on\" position.")
STEP(3, "Point the probes at a wall about 1 m (3 ft) away. In Home Assistant, turn on \"DIP stay awake\" and watch \"Distance to water\". It should read about 100 cm and change when you move.")
STEP(4, "Check \"Battery voltage\" reads 3.0-4.2 V, close to what your meter says.")
NOTE("Also worth checking if you have a meter that reads milliamps: with DIP asleep, the whole thing should draw well under 1 mA. "
     "Much more than that and the battery won't last.", "Optional: ")
P("If the distance stays empty or reads nonsense at 3.3V, stop here and ask for help. Don't move the sensor to a 5V pad on your own: "
  "at 5V its Echo signal is too strong for the SuperMini and would need two extra resistors.", italic=True)
BREAK()

# ================================================================ 9 probe cables
H("9. The probe cables")
P("The probes' plugs are too big to fit through the cable glands, so each cable gets cut, fed through, and soldered back together "
  "inside the box.")
IMG("probe_plug_gland.jpg", 3.6, "The plug's molded boot is bigger than the gland's hole.")
STEP(1, "Put a piece of tape on each probe cable right next to its plug and label them L and R, so the plug goes back on its own cable.")
STEP(2, "Cut each cable about 8 cm (3\") behind its plug.")
STEP(3, "Feed the long (probe) end of each cable in through a gland, from outside the box to inside.")
STEP(4, "Slide on heat shrink: one small piece over the centre wire, one larger piece that will cover the whole splice.")
STEP(5, "Strip both ends: cut back the black jacket about 15 mm, twist the braided shield into a tail, then strip 5 mm off the centre wire's insulation.")
STEP(6, "Solder centre to centre. Slide the small shrink over and heat it.")
STEP(7, "Solder shield to shield, laid alongside the first joint. Slide the big shrink over everything and heat it.")
IMG("d_coax.png", 6.2)
NOTE("Use the multimeter in beep mode to check: plug tip to the probe's centre should beep, plug tip to the shield should NOT.", "Check: ")

# ================================================================ 10 cable + cap
H("10. The 4-core cable and the pipe cap")
P("One 4-wire cable runs from the box to the pipe cap. Its four wires carry 3.3V, ground, the temperature signal and the float signal. "
  "Write down which color you use for each, here:")
TABLE([["Wire", "Job", "Your cable's color"],
       ["1", "3.3V", ""], ["2", "GND (ground)", ""], ["3", "Temp data", ""], ["4", "Float", ""]], [0.8, 2.4, 2.4])
H("10a. Box end", 2)
STEP(1, "Feed the cable through the third gland into the box. Strip back about 8 cm of the outer jacket.")
STEP(2, "Add the two pull-up resistors as shown below: the 4.7k from wire 1 to wire 3, and 100k (C) from wire 1 to wire 4. Solder each leg onto the wire with a little heat shrink over each joint.")
IMG("d_boxend.png", 6.2)
STEP(3, "Wire 1 to the holder's 3V pair 3 + (square). Wire 2 to 3V pair 3 - (round).")
STEP(4, "Wire 3 to SuperMini hole 5. Wire 4 to SuperMini hole 3.")
H("10b. Pipe cap end", 2)
STEP(1, "Feed the temperature probe's cable and the float wires up through the open part of the pipe cap's plate, and the 4-core cable in through the slot in the side of the cap.")
STEP(2, "Splice them as below. Use adhesive-lined heat shrink here: it's outdoors.")
IMG("d_capend.png", 6.2)
STEP(3, "Push the two ultrasonic probes up into the round holes in the cap from underneath. They should be a snug push fit; a small dab of silicone holds them if they're loose. Bring their cables out through the side slot too.")
STEP(4, "Screw the pipe cap top on with its 3 M3 screws.")
BREAK()

# ================================================================ 11 box together
H("11. Put the box together")
STEP(1, "Screw the battery holder onto its 4 round pads with its USB-C socket facing into the box.")
STEP(2, "Screw the sensor board onto its 4 posts with the 4 M2.5 x 6 screws, parts side down (connectors and pins facing you).")
STEP(3, "Slide the SuperMini's USB end under the two little lips, lay it flat, and screw the clamp tab over the other end with one M2.5 x 6 screw.")
IMG("r_esp_mount.png", 4.0, "SuperMini: USB end under the two lips, clamp tab over the far end")
STEP(4, "Plug the probe plugs into the sensor board. Tuck the wires so nothing is pinched.")
STEP(5, "Gently pull the cables back through the glands until they're snug, then tighten the gland caps by hand plus a quarter turn.")
STEP(6, "Put the battery in, switch to stay on. Put the lid on with the 4 M3 screws and rubber washers; tighten evenly in a cross pattern.")
NOTE("To charge later: take the lid off and plug a USB-C cable into the holder, or take the cell out and charge it in a charger.", "Charging: ")

# ================================================================ 12 software
H("12. Software: ESPHome and Home Assistant")
P("DIP runs ESPHome, which connects to Home Assistant. The project's firmware file is firmware/dip.yaml and the Home Assistant "
  "settings are in ha/dip_package.yaml.")
H("12a. Load the firmware (first time, over USB)", 2)
STEP(1, "Copy firmware/secrets.example.yaml to secrets.yaml and fill in the Wi-Fi name, password and an API key.")
STEP(2, "Connect the SuperMini to the computer with a USB-C cable (battery out).")
STEP(3, "Run: esphome run firmware/dip.yaml   (or use the ESPHome Builder app in Home Assistant).")
STEP(4, "If the board isn't found, hold the BOOT button, tap RESET, then let go of BOOT, and try again.")
P("After this first time, updates go over Wi-Fi. DIP sleeps most of the time, so first turn on \"DIP stay awake\" in "
  "Home Assistant and wait for its next wake (up to 20 minutes): it sees the switch and stays on. Turn the switch "
  "off when you're done. Left on, it drains the battery in a day or two.")
H("12b. Add it to Home Assistant", 2)
STEP(1, "Home Assistant usually finds DIP by itself: Settings > Devices & services, then Configure on the ESPHome device. Enter the API key.")
STEP(2, "Copy ha/dip_package.yaml into Home Assistant's packages folder. The steps at the top of the file show what to change: "
        "your two pump switches (the one that stops first, at 3\" down, and the one that stops at the float) and your phone.")
STEP(3, "Check the configuration (Developer tools > YAML > Check configuration), then restart Home Assistant.")
H("12c. Calibrate", 2)
STEP(1, "Level: with the pond full, read DIP's \"Distance to water\" and type that number into \"Pond full distance\". "
        "From then on \"Pond level below full\" shows how many inches the water has dropped.")
STEP(2, "Battery: measure the cell with the multimeter and compare it with DIP's \"Battery voltage\". If they're more than "
        "about 0.03 V apart, multiply the multiply: number under the battery sensor in dip.yaml by (meter reading / DIP's reading) "
        "and load the firmware again. A full cell then shows 100 %.")

# ================================================================ 13 install
H("13. Install at the pond")
STEP(1, "Drill 3-4 holes of 1/4\" around the bottom 2-3\" of the PVC pipe so water flows in and out, and one 1/8\" vent hole near the top.")
STEP(2, "Decide how low is \"too low\": below this the pumps all shut off. 6-8\" below the full level suits most ponds. "
        "Screw the float bracket inside the pipe at that height with the 2 M4 screws through the pipe wall (drill the pipe first).")
STEP(3, "Hang the temperature probe inside the pipe near the bottom.")
STEP(4, "Stand the pipe in the cinder block's hole in the pond. Pack pea gravel around it, checking it's straight up and down as you go.")
STEP(5, "Slide the pipe cap on, slot pointing toward the post. Pilot-drill the PVC through the hole in the cap's skirt and fit the short stainless screw so it can't turn or lift.")
STEP(6, "Hang the box on the post, glands at the bottom, with 2 stainless wood screws through the tab at the top. "
        "Run the cables along the back of the coping, under the net.")
NOTE("Cables come out of the cap sideways and go into the box from below, so rain runs off them instead of into anything.", "Why this way: ")

# ================================================================ 14 final checks
H("14. Final checks and troubleshooting")
H("Once it's installed", 2)
B("\"Distance to water\" and the pond level look right, and match your calibration.")
B("\"Water temperature\" reads a sensible number.")
B("Push the float down with a stick and hold it for a minute: \"Float switch low\" turns on within seconds, the pumps turn off and you get the alert. "
  "Let it go: the float clears, and the pumps come back on by themselves (the last one after 10 minutes, the first one at its next reading).")
B("\"Battery\" shows a percentage. Check it again in a week to see how fast it drops.")
H("If something's wrong", 2)
TABLE([["What you see", "Likely cause and fix"],
       ["Nothing shows up in Home Assistant", "Battery flat, switch not on \"stay on\", or Wi-Fi out of range. Check the SuperMini's red power light with the lid off."],
       ["Distance empty or jumping around", "Probe plugs loose, a probe splice broken, or a probe not pointing straight down. Check the splices with the meter (chapter 9)."],
       ["Temperature shows -127 or nothing", "The temp wire (3) isn't reaching pin 5, or the 4.7k is missing or broken."],
       ["Float always \"low\"", "Float upside down (chapter 6) or wire 4 touching ground somewhere."],
       ["Float never triggers", "Float stuck, mounted too low, or 100k (C) missing."],
       ["Battery reading 0 or wild", "Resistor (A)'s wire not on the + battery contact, or (A)/(B) not soldered in hole 4."],
       ["Battery drains in days", "It isn't going to sleep: make sure \"DIP stay awake\" is off in Home Assistant."]],
      [2.3, 4.3])
BREAK()

# ================================================================ appendix
H("Appendix: every connection in one table")
TABLE([["From", "To", "Notes"],
       ["Holder 3V pair 1 +", "SuperMini 3.3", "red"],
       ["Holder 3V pair 1 -", "SuperMini G", "black"],
       ["Holder + battery contact", "Resistor (A)", "via the purple wire"],
       ["Resistor (A) other leg", "SuperMini 4", "shares hole 4 with (B)"],
       ["Resistor (B)", "SuperMini 4 to SuperMini G", "legs sleeved, passes over 3.3"],
       ["Holder 3V pair 2 +", "Sensor board 5V pin", "female jumper"],
       ["Holder 3V pair 2 -", "Sensor board GND pin", "female jumper"],
       ["Sensor board Trig", "SuperMini 6", ""],
       ["Sensor board Echo", "SuperMini 7", ""],
       ["Probes (2)", "Sensor board white sockets", "cut and re-soldered inside the box"],
       ["Cable wire 1 (3.3V)", "Holder 3V pair 3 +", "4.7k and 100k (C) attach here"],
       ["Cable wire 2 (GND)", "Holder 3V pair 3 -", ""],
       ["Cable wire 3 (Temp)", "SuperMini 5", "4.7k from wire 1"],
       ["Cable wire 4 (Float)", "SuperMini 3", "100k (C) from wire 1"],
       ["Cap end: wire 1", "Temp probe red", ""],
       ["Cap end: wire 2", "Temp probe black + one float wire", "3-way splice"],
       ["Cap end: wire 3", "Temp probe yellow", ""],
       ["Cap end: wire 4", "Other float wire", ""]],
      [2.3, 2.6, 1.7])

doc.save(OUT)
print("saved", os.path.abspath(OUT))
